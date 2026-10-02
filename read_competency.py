import docx

def read_docx(path):
    print(f"--- READING: {path} ---")
    try:
        doc = docx.Document(path)
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                print(text)
        for t in doc.tables:
            print("[TABLE]")
            for row in t.rows:
                row_data = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                print(" | ".join(row_data))
    except Exception as e:
        print("Error reading:", e)
    print("\n\n")

read_docx('/Users/fahimasghar/Documents/QC-Lab-Manager/ISO 17025/Resource Requirement Procedure/(QCL-LSP-02) Personal Management/Training Management (6.2) 2025/2.09. Evaluation of Competency (FRM-2.09).docx')
